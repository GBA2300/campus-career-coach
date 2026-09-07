#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_resume_docx.py —— 由 JSON 生成排版好的中文简历 .docx（一页优先）。

用法:
    python build_resume_docx.py --data 简历.json --out "输出/张三-新媒体运营-简历v1.docx"
                                [--photo "证件照.jpg"] [--style classic|modern|compact]

依赖 python-docx；缺失时会自动尝试 pip 安装到当前 Python 环境。
字段说明见 assets/resume_schema.json，示例见 assets/resume_data.example.json。
"""
import argparse
import json
import os
import re
import subprocess
import sys

try:
    import docx
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("[提示] 未检测到 python-docx，正在安装…")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "python-docx"])
    import docx
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor


# ---------------- 样式配置 ----------------
STYLES = {
    "classic": {
        "name_size": 20, "name_font": "黑体", "name_bold": True,
        "title_size": 11.5, "title_font": "黑体", "title_color": "000000",
        "body_size": 10.5, "body_font": "宋体", "accent": "000000",
        "sub_size": 9.5, "line": 1.15, "margin_cm": 1.8, "bullet": "· ",
    },
    "modern": {
        "name_size": 22, "name_font": "黑体", "name_bold": True,
        "title_size": 11, "title_font": "黑体", "title_color": "1F4E79",
        "body_size": 10.5, "body_font": "宋体", "accent": "1F4E79",
        "sub_size": 9.5, "line": 1.15, "margin_cm": 1.7, "bullet": "▪ ",
    },
    "compact": {
        "name_size": 18, "name_font": "黑体", "name_bold": True,
        "title_size": 10.5, "title_font": "黑体", "title_color": "333333",
        "body_size": 10, "body_font": "宋体", "accent": "333333",
        "sub_size": 9, "line": 1.05, "margin_cm": 1.4, "bullet": "· ",
    },
}


# ---------------- 基础工具 ----------------
def set_run(run, cfg, size=None, bold=False, cn_font=None, color=None):
    f = run.font
    f.size = Pt(size or cfg["body_size"])
    f.bold = bold
    f.name = cn_font or cfg["body_font"]
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    fam = cn_font or cfg["body_font"]
    rf.set(qn("w:ascii"), fam)
    rf.set(qn("w:hAnsi"), fam)
    rf.set(qn("w:eastAsia"), fam)
    if color:
        f.color.rgb = RGBColor.from_string(color)
    return run


def new_para(doc, cfg, space_before=0, space_after=2, line=None, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line or cfg["line"]
    if align is not None:
        p.alignment = align
    return p


def add_bottom_border(p, color="888888", size=6):
    ppr = p._p.get_or_add_pPr()
    bdr = ppr.makeelement(qn("w:pBdr"), {})
    bottom = bdr.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    ppr.append(bdr)


def remove_table_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.makeelement(qn(f"w:{edge}"), {})
        el.set(qn("w:val"), "none")
        borders.append(el)
    pr.append(borders)


def setup_page(doc, cfg):
    for s in doc.sections:
        s.top_margin = Cm(cfg["margin_cm"])
        s.bottom_margin = Cm(cfg["margin_cm"])
        s.left_margin = Cm(cfg["margin_cm"])
        s.right_margin = Cm(cfg["margin_cm"])
    normal = doc.styles["Normal"]
    normal.font.name = cfg["body_font"]
    normal.font.size = Pt(cfg["body_size"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), cfg["body_font"])


# ---------------- 内容构建 ----------------
def build_header(doc, data, cfg, photo):
    b = data.get("basics", {})
    name = b.get("name", "").strip()
    if not name:
        raise ValueError("JSON 缺少 basics.name（简历必须有姓名）")

    contacts = []
    for key in ("phone", "email", "city", "political", "birth", "intention"):
        v = b.get(key)
        if v:
            contacts.append(str(v))
    for extra in b.get("extras", []) or []:
        if extra:
            contacts.append(str(extra))

    if photo and not os.path.exists(photo):
        print(f"[警告] 照片不存在，已跳过：{photo}")
        photo = None
    use_photo = bool(photo)

    if use_photo:
        t = doc.add_table(rows=1, cols=2)
        remove_table_borders(t)
        left, right = t.cell(0, 0), t.cell(0, 1)
        left.width = Cm(13.0)
        right.width = Cm(3.6)
        p_name = left.paragraphs[0]
    else:
        left = right = None
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(1)
    p_name.paragraph_format.line_spacing = cfg["line"]
    set_run(p_name.add_run(name), cfg, size=cfg["name_size"], bold=cfg["name_bold"],
            cn_font=cfg["name_font"], color=cfg["accent"])

    def add_line(text, size=None, bold=False, color=None):
        p = left.add_paragraph() if use_photo else doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = cfg["line"]
        if not use_photo:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), cfg, size=size or cfg["sub_size"], bold=bold, color=color)
        return p

    headline = b.get("headline") or b.get("intention")
    if headline and headline not in contacts:
        add_line(str(headline), size=cfg["body_size"], bold=True, color=cfg["accent"])
    if contacts:
        add_line("  |  ".join(contacts), size=cfg["sub_size"])
    links = b.get("links") or []
    if links:
        add_line("  |  ".join(f"{l.get('label', '链接')}：{l.get('url', '')}" for l in links),
                 size=cfg["sub_size"], color="555555")

    if use_photo:
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        try:
            rp.add_run().add_picture(photo, width=Cm(2.6))
        except Exception as e:
            print(f"[警告] 插入照片失败：{e}")


def build_section(doc, sec, cfg):
    title = sec.get("title", "").strip()
    stype = (sec.get("type") or "timeline").lower()
    items = sec.get("items") or []
    if not items:
        return

    if title:
        p = new_para(doc, cfg, space_before=8, space_after=4)
        r = p.add_run(title)
        set_run(r, cfg, size=cfg["title_size"], bold=True,
                cn_font=cfg["title_font"], color=cfg["title_color"])
        add_bottom_border(p, color=cfg["accent"])

    if stype in ("timeline", "exp", "edu", "project"):
        for it in items:
            _timeline(doc, it, cfg)
    elif stype in ("bullets", "list"):
        for it in items:
            text = it if isinstance(it, str) else (it.get("text") or "")
            p = new_para(doc, cfg, space_after=2)
            r = p.add_run(f"{cfg['bullet']}{text}")
            set_run(r, cfg)
    elif stype in ("text", "paragraph"):
        for it in items:
            text = it if isinstance(it, str) else (it.get("text") or "")
            p = new_para(doc, cfg, space_after=3)
            r = p.add_run(text)
            set_run(r, cfg)
    elif stype == "kv":
        for it in items:
            p = new_para(doc, cfg, space_after=2)
            k = it.get("k") or it.get("key") or ""
            v = it.get("v") or it.get("value") or ""
            r1 = p.add_run(f"{cfg['bullet']}{k}：")
            set_run(r1, cfg, bold=True)
            r2 = p.add_run(str(v))
            set_run(r2, cfg)
    else:
        for it in items:
            text = it if isinstance(it, str) else json.dumps(it, ensure_ascii=False)
            p = new_para(doc, cfg, space_after=2)
            set_run(p.add_run(f"{cfg['bullet']}{text}"), cfg)


def _timeline(doc, it, cfg):
    org = (it.get("org") or it.get("company") or it.get("school") or "").strip()
    role = (it.get("role") or it.get("title") or "").strip()
    period = (it.get("period") or it.get("date") or "").strip()
    loc = (it.get("location") or "").strip()

    p = new_para(doc, cfg, space_before=3, space_after=1)
    left_parts = [x for x in (org, role) if x]
    if left_parts:
        r = p.add_run("  |  ".join(left_parts))
        set_run(r, cfg, size=cfg["body_size"], bold=True)
    right_parts = [x for x in (loc, period) if x]
    if right_parts:
        r2 = p.add_run("\t" + "  |  ".join(right_parts))
        set_run(r2, cfg, size=cfg["sub_size"], color="555555")
        _right_tab(p, doc, cfg)

    for b in it.get("bullets", []) or []:
        if not b:
            continue
        bp = new_para(doc, cfg, space_after=1)
        bp.paragraph_format.left_indent = Cm(0.45)
        set_run(bp.add_run(f"{cfg['bullet']}{b}"), cfg)


def _right_tab(p, doc, cfg):
    from docx.shared import Cm as _Cm
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(width, 2)  # 2 = RIGHT


# ---------------- 主流程 ----------------
def build(data, out, photo=None, style="classic"):
    cfg = STYLES.get(style, STYLES["classic"])
    doc = Document()
    setup_page(doc, cfg)
    photo = photo or data.get("basics", {}).get("photo")
    build_header(doc, data, cfg, photo)
    for sec in data.get("sections", []) or []:
        build_section(doc, sec, cfg)

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    doc.save(out)

    n_items = sum(len(s.get("items") or []) for s in data.get("sections", []))
    n_bul = sum(len((i.get("bullets") or [])) for s in data.get("sections", [])
                for i in (s.get("items") or []) if isinstance(i, dict))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    print("=" * 56)
    print(f"已生成：{out}")
    print(f"  风格 {style} | 板块 {len(data.get('sections', []))} 个 | 条目 {n_items} 条 | 经历 bullet {n_bul} 条")
    print(f"  正文字数（中文）约 {len(re.findall(r'[一-鿿]', text))} 字")
    if photo and os.path.exists(photo):
        print(f"  已插入照片：{photo}")
    print("=" * 56)
    print("下一步：用 present_files 打开预览，并附「改了什么 / 为什么改」说明。")


def main():
    ap = argparse.ArgumentParser(description="JSON → 中文简历 .docx")
    ap.add_argument("--data", required=True, help="简历数据 JSON 路径")
    ap.add_argument("--out", required=True, help="输出 .docx 路径")
    ap.add_argument("--photo", help="证件照路径（jpg/png）")
    ap.add_argument("--style", default="classic", choices=list(STYLES.keys()))
    args = ap.parse_args()

    if not os.path.exists(args.data):
        print(f"[错误] JSON 不存在: {args.data}", file=sys.stderr)
        sys.exit(1)
    with open(args.data, "r", encoding="utf-8") as f:
        data = json.load(f)

    try:
        build(data, args.out, args.photo, args.style)
    except Exception as e:
        print(f"[错误] 生成失败: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
